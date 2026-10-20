# -*- coding: utf-8 -*-
"""בונה tests/template/fixture.pptx — תסריט סינתטי קטן עם פגמים מושתלים ומלכודות
לבדיקות #37-#45 של template_scan.py. הרצה: python build_fixture.py [out.pptx]"""
import io
import os
import sys

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

W, H = 12192000, 6858000
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(__file__), "fixture.pptx")

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
blank = prs.slide_layouts[6]


def tb(slide, x, y, w, h, paras, size=20, color=None, name=None, fill=None):
    s = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    if name:
        s.name = name
    tf = s.text_frame
    tf.word_wrap = True
    for i, t in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        for j, part in enumerate(t.split("")):
            if j:
                p.add_line_break()
            r = p.add_run()
            r.text = part
            r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = RGBColor.from_string(fill)
    return s


def radio(slide, x, y):
    e = slide.shapes.add_shape(MSO_SHAPE.OVAL, Emu(x), Emu(y), Emu(252766), Emu(261512))
    e.name = "אליפסה"
    return e


def option_rows(slide, texts, y0=3200000):
    for i, t in enumerate(texts):
        y = y0 + i * 750000
        tb(slide, 5265197, y, 5741155, 400110, [t], name="מסיח %d" % (i + 1))
        radio(slide, 11087846, y + 60000)


def feedback_boxes(slide, n=2):
    for i in range(n):
        s = tb(slide, -8825913 + i * 4300000, 4085188, 4167984, 2000000,
               ["משוב %s: ההסבר המלא מופיע כאן ומנמק את התשובה." % ("להצלחה" if i == 0 else "לאי הצלחה")])
        s.name = "תיבת משוב %d" % (i + 1)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def ui(slide):
    tb(slide, 268356, 6182143, 1798983, 506891, ["צדקתי?"])


# ---- שקף 1: ריבוי-תשובות, 4 מסיחים. נטועים: #37 #2 #44 #40(לבן) #45 ; תקין: 2 משובים
s = prs.slides.add_slide(blank)
tb(s, 3831910, 1156726, 7735614, 1900000,
   ["בניסוי מדדו את המסה של שלושה גופים על מאזניים.", "אילו מהמשפטים הבאים נכונים?", "לחצו על כל התשובות הנכונות."])
option_rows(s, ["המסה תלויה במיקום.", "המסה היא כמות החומר בגוף", "משקל   שונה ממסה", "המסה נמדדת בק״ג"])
feedback_boxes(s, 2)
tb(s, 624476, 5800000, 3000000, 400000, ["טקסט נסתר בלבן"], color="FFFFFF")
ui(s)
notes(s, "הנחיות הפקה:\nשם המסך בפיגמה: MultipleChoiceQuestion\nהתשובה הנכונה - ב, ד")

# ---- שקף 2: 2 מסיחים עם רמז. נטועים: #38 #42(0 משובים) #40(אנימציה+למפתחת) #41(תמונה-טבלה)
s = prs.slides.add_slide(blank)
tb(s, 3831910, 1156726, 7735614, 1200000, ["לפי הנתונים בטבלה, איזה גוף כבד יותר?"])
option_rows(s, ["הכדור", "הקובייה"])
tb(s, 268356, 5600000, 1798983, 506891, ["   אפשר רמז?"])
ui(s)
tb(s, 624476, 1445965, 3880694, 600164, ["אנימציה: הכפתור מהבהב בכניסה"])
tb(s, 624476, 2187778, 3880694, 769441, ["למפתחת:", "צייני מטה בהערות הפקה האם נדרש סרגל"])
buf = io.BytesIO()
Image.new("RGB", (400, 300), "white").save(buf, "PNG")
buf.seek(0)
pic = s.shapes.add_picture(buf, Emu(700000), Emu(3300000), Emu(3000000), Emu(2200000))
pic.name = "טבלה_נתונים.png"
notes(s, "שם המסך בפיגמה: SingleChoiceQuestion")

# ---- שקף 3: מלכודות — הכול תקין
s = prs.slides.add_slide(blank)
tb(s, 3831910, 1156726, 7735614, 1500000,
   ["כדור ברזל וכדור עץ בגודל זהה הונחו על מאזניים.", "", "איזה כדור מכיל יותר חומר?"])
option_rows(s, ["כדור הברזל", "כדור העץ"])
feedback_boxes(s, 2)
dark = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(268356), Emu(6182143), Emu(1798983), Emu(506891))
dark.fill.solid()
dark.fill.fore_color.rgb = RGBColor(0x01, 0x9D, 0xE5)
dark.text_frame.text = "בדקו אותי"
dark.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
tb(s, 5265197, 5600000, 3000000, 400000, ["שאלה 1 מתוך 5"], size=11)
tb(s, 624476, 1445965, 3000000, 500000, ["מקום לגיף"])
tb(s, 624476, 2187778, 3000000, 500000, ["איריס, כדאי לקצר את השאלה הזו"])
tbl = s.shapes.add_table(2, 2, Emu(700000), Emu(3000000), Emu(3000000), Emu(800000))
tbl.table.cell(0, 0).text = "גוף"
tbl.table.cell(0, 1).text = "מסה"
tbl.table.cell(1, 0).text = "כדור"
tbl.table.cell(1, 1).text = "50 גרם"
ui(s)
notes(s, "שם המסך בפיגמה: SingleChoiceQuestion\nשאלה אוריינית: לא")

# ---- שקף 4: אותו גזע כמו שקף 3 (שקף-מצב נפרד) + מילה בודדת. נטועים: #42(גזע זהה) #43
s = prs.slides.add_slide(blank)
tb(s, 3831910, 1156726, 7735614, 1500000,
   ["כדור ברזל וכדור עץ בגודל זהה הונחו על מאזניים.", "", "איזה כדור מכיל יותר חומר?"])
option_rows(s, ["כדור הברזל", "כדור העץ"])
feedback_boxes(s, 2)
tb(s, 624476, 4000000, 3000000, 1200000,
   ["המאזניים נשארו מאוזנים גם כשהחלפנו את הגופים\vכן"])
ui(s)
notes(s, "שם המסך בפיגמה: SingleChoiceQuestion")

prs.save(OUT)
print("saved", OUT)
