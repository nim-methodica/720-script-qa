# -*- coding: utf-8 -*-
"""template_scan.py — סריקה דטרמיניסטית ברמת-ה-PPTX לבדיקות #37-#45 (תקני תבנית וניקיון-הפקה)
+ הרחבת #2 (רווח כפול) לכל תיבות-הטקסט, התאים והמסיחים.

מדוע על ה-PPTX ולא על הטקסט המחולץ: markitdown משטח את השקף — לא רואים אילו תיבות הן מסיחים,
מה מחוץ לקנבס (משובים/רמז חונים שם), באיזה צבע הטקסט, האם "טבלה" היא טבלה חיה או תמונה, ואיפה
נשברות השורות. הסקריפט קורא את ה-XML ישירות (כולל shapes העטופים ב-mc:AlternateContent).

בדיקות (מספור לפי references/qa-dimensions.md, קבוצה ח):
  #37 נקודה בסוף מסיח            #38 שאלה בת 2 מסיחים — האם יש רמז / ניסיון שני
  #39 סימון שאלות ריבוי-תשובות (לקריאה שיפוטית של המשוב)
  #40 ניקיון קנבס: טקסט לבן / אלמנט מוסתר / שרידי-תבנית והנחיות / אנימציה כתובה על המסך
  #41 טקסט חי: תמונות במקום טקסט/טבלה/נוסחה  #42 ספירת תיבות-משוב (צריך 2) + שקפי-מצב נפרדים
  #43 מילה בודדת בשורה (שבירה ידנית / הערכת גלישה) #44 רווח בין הטקסט המקדים לשאלה
  #45 תשובה נכונה כתובה בהערות-דובר   +  #2 רווח כפול בכל מקום

חשוב: הסקריפט מדווח **מועמדים** עם ציטוט — הכרעה אחרונה (ובעיקר #39, #41, #43 "גבולי") היא שיפוטית.

שימוש:
  PYTHONUTF8=1 python template_scan.py "<file.pptx>"
  PYTHONUTF8=1 python template_scan.py "<file.pptx>" --slides 28,30
"""
import io
import os
import re
import sys
from collections import Counter
from difflib import SequenceMatcher

from lxml import etree
from pptx import Presentation

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
}
A = "{%s}" % NS["a"]
P = "{%s}" % NS["p"]
MC = "{%s}" % NS["mc"]

EMU_PT = 12700.0
UI_LABELS = ("צדקתי", "אפשר רמז", "חזרה", "המשך", "חזרה לשאלה", "בחרתי", "קחו אותי")
IMPERATIVE_OPENERS = (
    "לחצו", "בחרו", "סמנו", "גררו", "הזינו", "כתבו", "חשבו", "מצאו", "השלימו", "התאימו", "סדרו",
    "מיינו", "קבעו", "הסבירו", "איזה", "איזו", "מה", "מי", "כמה", "האם", "מדוע", "למה", "איך", "היכן",
    "מתי", "אילו", "באיזה", "באיזו", "על איזה", "מהו", "מהי", "מהם", "מהן",
)
SCREEN_RE = re.compile(r"שם המסך[^:\n：]*[:：]\s*([A-Za-z][A-Za-z \-]*)")
QUESTION_TOKEN_RE = re.compile(r"Question|Choice|Hotspot|Dropdown|DragAndDrop|ValueInput|Statement", re.I)
MULTI_TOKEN_RE = re.compile(r"MultipleChoice|MultiSelect|MultipleSelection", re.I)
MULTI_TEXT_RE = re.compile(r"כל (?:ה)?תשובות|את כל ה|יותר מתשובה אחת|כל המשפטים הנכונים|כל ה[א-ת]+ הנכונ")
FEEDBACK_OPENER_RE = re.compile(r"^\s*(נכון|לא נכון|כל הכבוד|טעית|כמעט|מצוין|יפה|עוד לא|שימו לב|אופס|בדיוק|נכון מאוד)")
LEFTOVER_RES = [
    (re.compile(r"למפתח(?:ת|/ת)?\s*[:\-–]"), "הנחיה 'למפתח/ת' על המסך"),
    (re.compile(r"כאן (?:ה)?הנחיה|כאן שמים"), "טקסט-הנחיה של התבנית"),
    (re.compile(r"טקסט H1|טקסט רץ|טקסט השאלה|טקסט רמז"), "טקסט-דמה של התבנית"),
    (re.compile(r"(?<![א-ת])מסיח \d(?!\d)"), "'מסיח N' — פלייסהולדר"),
    (re.compile(r"משוב להצלחה \(|משוב לאי הצלחה"), "תווית-תבנית של תיבת-משוב"),
    (re.compile(r"הנחיות הפקה|שם המסך בפ"), "'הנחיות הפקה'/'שם המסך' על הקנבס (שייך להערות-דובר)"),
]
ANIMATION_RE = re.compile(
    r"אנימציה|אנימצייה|אנימט|אפקט|Fade|Appear|Wipe|בלחיצה על|ייכנס|תיכנס|נכנס(?:ת)? מ|מופיע(?:ה)? בהדרגה|יופיע בהדרגה"
)
ANSWER_IN_NOTES_RE = re.compile(
    r"(?:ה)?תשובה (?:ה)?נכונה|(?:ה)?תשובות (?:ה)?נכונות|מפתח(?: ה)?תשובות|התשובה (?:היא|הנכונה)|תשובה:\s*\S"
)
PERSONAL_NOTE_RE = re.compile(r"^\s*(איריס|אפרת)\b")      # הערות-ביקורת בשם — מחוץ להיקף (SKILL.md עיקרון #6)
TABLE_WORD_RE = re.compile(r"טבלה|בטבלה|הטבלה")
PIC_KEYWORD_RE = re.compile(r"טבלה|table|נוסחה|formula|משוואה|equation|תרגיל|screenshot|צילום[ -]?מסך|רמז", re.I)


# ----------------------------------------------------------------- fonts / measuring
_FONT_DIRS = [
    os.path.join(os.environ.get("LOCALAPPDATA", ""), "Microsoft", "Windows", "Fonts"),
    r"C:\Windows\Fonts", os.path.expanduser("~/.fonts"), "/usr/share/fonts", "/Library/Fonts",
]
_font_cache = {}
FONT_USED = {"name": None}


def _find_font(bold):
    cands = (["Assistant-Bold.ttf", "Assistant-SemiBold.ttf"] if bold else ["Assistant-Regular.ttf"]) + \
            (["arialbd.ttf"] if bold else ["arial.ttf"])
    for fn in cands:
        for d in _FONT_DIRS:
            pth = os.path.join(d, fn)
            if os.path.exists(pth):
                return pth, fn
    return None, None


def measure(text, size_pt, bold=False):
    """רוחב-טקסט בנקודות (pt). None אם אין PIL/גופן."""
    try:
        from PIL import ImageFont
    except Exception:
        return None
    key = bold
    if key not in _font_cache:
        pth, fn = _find_font(bold)
        if not pth:
            _font_cache[key] = None
        else:
            _font_cache[key] = pth
            FONT_USED["name"] = FONT_USED["name"] or fn
    pth = _font_cache[key]
    if not pth:
        return None
    k2 = (pth, round(size_pt * 4))
    if k2 not in _font_cache:
        _font_cache[k2] = ImageFont.truetype(pth, max(1, round(size_pt * 4)))
    return _font_cache[k2].getlength(text) / 4.0


# ----------------------------------------------------------------- XML model
class Node:
    __slots__ = ("tag", "name", "descr", "hidden", "box", "paras", "fill", "prst", "depth",
                 "group", "el", "is_pic", "is_table", "wrap", "ins", "autofit")

    def __init__(self):
        self.box = None
        self.paras = []
        self.fill = None
        self.prst = None
        self.group = None
        self.hidden = False
        self.is_pic = False
        self.is_table = False
        self.descr = ""
        self.wrap = "square"
        self.ins = (91440, 45720, 91440, 45720)
        self.autofit = None

    @property
    def text(self):
        return "\n".join(p["text"] for p in self.paras)

    @property
    def has_text(self):
        return any(p["text"].strip() for p in self.paras)


def _xfrm(el):
    for path in ("p:spPr/a:xfrm", "p:grpSpPr/a:xfrm", "p:xfrm", "p:pic/p:spPr/a:xfrm"):
        x = el.find(path, NS)
        if x is not None:
            return x
    return None


def _color_of(solid):
    if solid is None:
        return None
    c = solid.find("a:srgbClr", NS)
    if c is not None:
        return c.get("val").upper()
    c = solid.find("a:schemeClr", NS)
    if c is not None:
        v = c.get("val")
        if v in ("bg1", "lt1") and c.find("a:lumMod", NS) is None:
            return "FFFFFF"
        return "scheme:" + v
    return None


def _para_model(p):
    segs = []
    white_runs = []
    sz = None
    bold = False
    for ch in p:
        t = etree.QName(ch).localname
        if t in ("r", "fld"):
            tx = "".join(x.text or "" for x in ch.findall("a:t", NS))
            segs.append(tx)
            rpr = ch.find("a:rPr", NS)
            if rpr is not None:
                if sz is None and rpr.get("sz"):
                    sz = int(rpr.get("sz")) / 100.0
                if rpr.get("b") == "1":
                    bold = True
                col = _color_of(rpr.find("a:solidFill", NS))
                if col == "FFFFFF" and tx.strip():
                    white_runs.append(tx.strip())
        elif t == "br":
            segs.append("\v")
    ppr = p.find("a:pPr", NS)
    spc_bef = spc_aft = 0
    if ppr is not None:
        sb = ppr.find("a:spcBef/a:spcPts", NS)
        sa = ppr.find("a:spcAft/a:spcPts", NS)
        spc_bef = int(sb.get("val")) if sb is not None else 0
        spc_aft = int(sa.get("val")) if sa is not None else 0
        if ppr.find("a:spcBef/a:spcPct", NS) is not None:
            spc_bef = max(spc_bef, int(ppr.find("a:spcBef/a:spcPct", NS).get("val")) // 100)
        if ppr.find("a:spcAft/a:spcPct", NS) is not None:
            spc_aft = max(spc_aft, int(ppr.find("a:spcAft/a:spcPct", NS).get("val")) // 100)
    return {"text": "".join(segs), "sz": sz, "bold": bold, "white": white_runs,
            "spc_bef": spc_bef, "spc_aft": spc_aft}


def _node_from(el, tag, depth, group, tf):
    n = Node()
    n.tag, n.depth, n.group, n.el = tag, depth, group, el
    cnv = el.find(".//p:cNvPr", NS)
    if cnv is not None:
        n.name = cnv.get("name") or ""
        n.descr = cnv.get("descr") or ""
        n.hidden = cnv.get("hidden") == "1"
    else:
        n.name = ""
    xf = _xfrm(el)
    if xf is not None:
        off, ext = xf.find("a:off", NS), xf.find("a:ext", NS)
        if off is not None and ext is not None:
            x, y = int(off.get("x")), int(off.get("y"))
            cx, cy = int(ext.get("cx")), int(ext.get("cy"))
            sx, sy, tx, ty = tf
            n.box = (tx + sx * x, ty + sy * y, sx * cx, sy * cy)
    if tag == "pic":
        n.is_pic = True
    sppr = el.find("p:spPr", NS)
    if sppr is not None:
        if sppr.find("a:blipFill", NS) is not None:
            n.is_pic = True
        if sppr.find("a:noFill", NS) is not None:
            n.fill = "none"
        sf = sppr.find("a:solidFill", NS)
        if sf is not None:
            n.fill = _color_of(sf) or "solid?"
        if sppr.find("a:gradFill", NS) is not None:
            n.fill = "grad"
        g = sppr.find("a:prstGeom", NS)
        if g is not None:
            n.prst = g.get("prst")
    if n.fill is None and el.find("p:style", NS) is not None:
        n.fill = "style"
    body = el.find("p:txBody", NS)
    if body is not None:
        bp = body.find("a:bodyPr", NS)
        if bp is not None:
            n.wrap = bp.get("wrap", "square")
            n.ins = (int(bp.get("lIns", 91440)), int(bp.get("tIns", 45720)),
                     int(bp.get("rIns", 91440)), int(bp.get("bIns", 45720)))
            n.autofit = "norm" if bp.find("a:normAutofit", NS) is not None else None
        n.paras = [_para_model(p) for p in body.findall("a:p", NS)]
    if tag == "graphicFrame":
        tbl = el.find(".//a:tbl", NS)
        if tbl is not None:
            n.is_table = True
            for tc in tbl.iter(A + "tc"):
                tb = tc.find("a:txBody", NS)
                if tb is not None:
                    n.paras.extend(_para_model(p) for p in tb.findall("a:p", NS))
    return n


def iter_nodes(parent, depth=0, group=None, tf=(1.0, 1.0, 0.0, 0.0)):
    for ch in parent:
        t = etree.QName(ch).localname
        if t == "AlternateContent":
            choice = ch.find("mc:Choice", NS)
            src = choice if choice is not None else ch.find("mc:Fallback", NS)
            if src is not None:
                yield from iter_nodes(src, depth, group, tf)
        elif t in ("sp", "pic", "graphicFrame", "cxnSp"):
            yield _node_from(ch, t, depth, group, tf)
        elif t == "grpSp":
            g = _node_from(ch, "grpSp", depth, group, tf)
            yield g
            gx = ch.find("p:grpSpPr/a:xfrm", NS)
            ntf = tf
            if gx is not None:
                off, ext = gx.find("a:off", NS), gx.find("a:ext", NS)
                co, ce = gx.find("a:chOff", NS), gx.find("a:chExt", NS)
                if None not in (off, ext, co, ce) and int(ce.get("cx")) and int(ce.get("cy")):
                    kx = int(ext.get("cx")) / int(ce.get("cx"))
                    ky = int(ext.get("cy")) / int(ce.get("cy"))
                    ox, oy = int(off.get("x")), int(off.get("y"))
                    cox, coy = int(co.get("x")), int(co.get("y"))
                    sx, sy, tx, ty = tf
                    ntf = (sx * kx, sy * ky,
                           tx + sx * (ox - cox * kx), ty + sy * (oy - coy * ky))
            yield from iter_nodes(ch, depth + 1, g, ntf)


# ----------------------------------------------------------------- helpers
def off_canvas(box, W, H):
    if not box:
        return False
    x, y, w, h = box
    return x + w <= 0 or x >= W or y + h <= 0 or y >= H


def lum(hexrgb):
    r, g, b = (int(hexrgb[i:i + 2], 16) for i in (0, 2, 4))
    return (0.299 * r + 0.587 * g + 0.114 * b) / 255.0


def words(s):
    return [w for w in re.split(r"\s+", s.strip()) if re.search(r"[\w\u0590-\u05FF]", w)]


def snippet(s, n=48):
    s = s.replace("\v", " ⏎ ").replace("\n", " | ")
    return s if len(s) <= n else s[:n] + "…"


def is_ui_label(txt):
    t = txt.strip()
    return any(t.startswith(u) for u in UI_LABELS) and len(words(t)) <= 3


def wrap_lines(text, size_pt, bold, width_pt):
    """חלוקה-חמדנית למילים. מחזיר רשימת שורות (כל שורה = רשימת מילים) או None."""
    if measure("א", size_pt, bold) is None:
        return None
    space = measure(" ", size_pt, bold) or size_pt * 0.25
    lines, cur, cur_w = [], [], 0.0
    for w in text.split():
        ww = measure(w, size_pt, bold)
        add = ww if not cur else space + ww
        if cur and cur_w + add > width_pt:
            lines.append(cur)
            cur, cur_w = [w], ww
        else:
            cur.append(w)
            cur_w += add
    if cur:
        lines.append(cur)
    return lines


# ----------------------------------------------------------------- per-slide scan
def scan_slide(slide, idx, W, H, prev_stem):
    out = []   # (check, severity, detail)
    nodes = list(iter_nodes(slide.shapes._spTree))
    notes = ""
    try:
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text or ""
    except Exception:
        pass

    text_nodes = [n for n in nodes if n.has_text and n.tag in ("sp", "graphicFrame")]
    on = [n for n in text_nodes if not off_canvas(n.box, W, H)]
    off = [n for n in text_nodes if off_canvas(n.box, W, H)]
    ellipses = [n for n in nodes if n.tag == "sp" and n.prst == "ellipse" and n.box
                and n.box[2] < 600000 and n.box[3] < 600000 and not n.has_text]

    sm = SCREEN_RE.search(notes)
    screen = sm.group(1).strip() if sm else ""
    all_on_text = " ".join(n.text for n in on)
    is_question = bool(QUESTION_TOKEN_RE.search(screen)) or "צדקתי" in all_on_text or "אפשר רמז" in all_on_text
    is_multi = bool(MULTI_TOKEN_RE.search(screen)) or bool(MULTI_TEXT_RE.search(all_on_text))

    # ---- אפשרויות-בחירה (מסיחים): טקסט על-קנבס שלידו עיגול-רדיו קטן באותו גובה
    options = []
    if is_question:
        for n in on:
            if not n.box or is_ui_label(n.text) or n.tag != "sp":
                continue
            x, y, w, h = n.box
            if h > 1300000:      # גזע/פאנל גבוה — לא שורת-מסיח
                continue
            cy = y + h / 2.0
            for e in ellipses:
                ex, ey, ew, eh = e.box
                if abs((ey + eh / 2.0) - cy) <= max(eh * 1.2, h * 0.65) and (x - 700000) <= ex <= (x + w + 700000):
                    options.append(n)
                    break
        if not options:   # נפילה לשם-צורה
            options = [n for n in on if re.search(r"מסיח", n.name)]

    # ---- #37 נקודה בסוף מסיח
    for n in options:
        t = n.text.strip()
        if t.endswith((".", "。")) and not re.search(r"\d\.$", t):
            out.append((37, "🔵", f'מסיח מסתיים בנקודה: "{snippet(t)}"'))

    # ---- #38 שאלה בת 2 מסיחים: ניסיון אחד, ללא רמז
    hint_shapes = [n for n in nodes if (n.has_text and re.search(r"אפשר רמז|חזרה לשאלה", n.text))]
    hint_groups = [g for g in nodes if g.tag == "grpSp" and any(
        c.has_text and re.search(r"חזרה לשאלה|טקסט רמז|רמז", c.text) for c in nodes if c.group is g)]
    if is_question and len(options) == 2:
        evid = []
        if any(re.search(r"אפשר רמז", n.text) for n in hint_shapes):
            evid.append('כפתור "אפשר רמז?" על הקנבס')
        if hint_groups:
            evid.append("קבוצת-רמז (פאנל רמז) בשקף")
        if re.search(r"ניסיון שני|שני ניסיונות|2 ניסיונות|שתי הזדמנויות|רמז", notes):
            evid.append("הערות-דובר מזכירות רמז/ניסיון שני")
        if evid:
            out.append((38, "🟠", "שאלה עם 2 מסיחים: " + " · ".join(evid) +
                        " — צריך ניסיון מענה אחד, בלי רמז"))
        else:
            out.append(("info", "", "שאלה עם 2 מסיחים — לא נמצא רמז (תקין); לוודא שגם ההערות לא דורשות ניסיון שני"))

    # ---- #39 ריבוי-תשובות (סימון לקריאה שיפוטית)
    if is_question and is_multi:
        out.append(("info", "", "שאלת ריבוי-תשובות — לבדוק ידנית (#39) שהמשובים מנוסחים לריבוי-תשובות ומתאימים "
                                "גם למצב שבו מערכת דורסת את בחירת התלמיד/ה (תבנית שקף 16)"))

    # ---- #42 תיבות-משוב מחוץ לקנבס (צריך בדיוק 2) + שקפי-מצב נפרדים
    if is_question:
        hint_grp_ids = {id(g) for g in hint_groups}
        fb = []
        for n in off:
            if n.tag != "sp" or n.group is not None:
                continue
            if re.search(r"חזרה לשאלה|טקסט רמז", n.text) or n.text.lstrip().startswith("רמת חשיבה"):
                continue
            if not (re.search(r"משוב", n.text) or (n.box and n.box[3] >= 1200000)):
                continue          # פריט-גרירה/תווית חונים — לא תיבת-משוב
            fb.append(n)
        for n in on:   # משוב על הקנבס שנפתח בפתיח-משוב אופייני (לא מסיח/תווית)
            if n.tag == "sp" and n not in options and FEEDBACK_OPENER_RE.match(n.text) and n.group is None \
                    and not is_ui_label(n.text):
                fb.append(n)
        if len(fb) != 2:
            out.append((42, "🟠", f"נמצאו {len(fb)} תיבות-משוב בשקף השאלה (נדרשות 2: הצלחה / אי-הצלחה, "
                                  "מחוץ לקנבס משמאל)" +
                        ("" if not fb else " — " + " || ".join(f'"{snippet(n.text, 30)}"' for n in fb))))
        else:
            out.append(("info", "", "2 תיבות-משוב — תקין: " + " || ".join(f'"{snippet(n.text, 28)}"' for n in fb)))
    stem = ""
    cand = [n for n in on if n.tag == "sp" and n not in options and not is_ui_label(n.text)
            and len(words(n.text)) >= 6]
    if cand:
        stem = max(cand, key=lambda n: len(n.text)).text.strip()
    if is_question and stem and prev_stem and len(stem) > 25 \
            and SequenceMatcher(None, stem, prev_stem).ratio() >= 0.9:
        out.append((42, "🟠", f'גזע-השאלה זהה לשקף הקודם — נראה כשקף-מצב נפרד של אותה שאלה: "{snippet(stem, 40)}" '
                              "(בתבנית העדכנית: שקף אחד לכל שאלה עם שני משובים)"))

    # ---- #45 תשובה בהערות-דובר
    for ln in notes.splitlines():
        if ANSWER_IN_NOTES_RE.search(ln):
            out.append((45, "🟠", f'תשובה נכונה בהערות-דובר: "{snippet(ln.strip(), 60)}"'))

    # ---- #40 ניקיון קנבס
    fullbleed = any(n.is_pic and n.box and (n.box[2] * n.box[3]) >= 0.6 * W * H for n in nodes)
    filled = [n for n in nodes if n.tag == "sp" and n.fill and n.fill not in ("none",) and n.box
              and not (n.has_text and False)]
    for n in on:
        for p in n.paras:
            if not p["white"]:
                continue
            if fullbleed or n.fill not in (None, "none", "FFFFFF"):
                continue
            if n.group is not None and any(g.fill not in (None, "none", "FFFFFF") for g in [n.group]):
                continue
            covered = False
            if n.box:
                cx, cy = n.box[0] + n.box[2] / 2.0, n.box[1] + n.box[3] / 2.0
                for f in filled:
                    if f is n or f.fill == "FFFFFF":
                        continue
                    fx, fy, fw, fh = f.box
                    if fx <= cx <= fx + fw and fy <= cy <= fy + fh:
                        if f.fill and not f.fill.startswith(("scheme", "style", "grad", "solid")):
                            if lum(f.fill) >= 0.6:
                                continue
                        covered = True
                        break
            if not covered:
                out.append((40, "🟠", f'טקסט לבן ללא רקע כהה מאחוריו (בלתי-נראה/מיותר?): "{snippet("".join(p["white"]), 50)}"'))
    for n in nodes:
        if n.hidden:
            out.append((40, "🟠", f'אלמנט מוסתר (hidden) בשקף: "{n.name}" {snippet(n.text, 40)}'))
    for n in on:
        txt = n.text
        if PERSONAL_NOTE_RE.search(txt):
            continue   # הערת-ביקורת בשם — מחוץ להיקף (SKILL.md עיקרון #6)
        for rx, label in LEFTOVER_RES:
            m = rx.search(txt)
            if m:
                out.append((40, "🔵", f'{label}: "{snippet(txt, 50)}"'))
                break
        am = ANIMATION_RE.search(txt)
        if am and not any(c == 40 and "אנימציה" in d and snippet(txt, 40) in d for c, _, d in out):
            out.append((40, "🔵", f'תיאור אנימציה כתוב על המסך ("{am.group(0)}") — להעביר להערות-דובר: "{snippet(txt, 50)}"'))

    # ---- #41 טקסט חי
    tables = [n for n in nodes if n.is_table]
    pics = []
    for n in nodes:
        if not n.is_pic or not n.box:
            continue
        w, h = n.box[2], n.box[3]
        frac = (w * h) / float(W * H)
        if frac >= 0.85 or frac < 0.015:
            continue
        if n.box[1] < 0.08 * H and h < 0.14 * H:   # פס-מסע/לוגו בראש התבנית
            continue
        if off_canvas(n.box, W, H):
            continue
        pics.append((n, frac))
    for n, frac in pics:
        kw = PIC_KEYWORD_RE.search(n.descr + " " + n.name)
        if kw:
            out.append((41, "🟠", f'תמונה שנראית כטקסט/טבלה/נוסחה ("{kw.group(0)}" בשם/alt): '
                                  f'"{snippet(n.name + " | " + n.descr, 60)}" ({frac * 100:.0f}% מהקנבס) — להמיר לטקסט חי'))
    if pics and TABLE_WORD_RE.search(all_on_text) and not tables:
        out.append((41, "🟠", "הטקסט מזכיר טבלה אך אין בשקף טבלה חיה (יש תמונה/ות) — לבדוק בעין שהטבלה אינה תמונה"))
    if is_question and pics:
        out.append(("info", "", "תמונות בשקף שאלה לבדיקת-עין (#41: טקסט/נוסחה/טבלה בתוך תמונה?): " +
                    "; ".join(f'{n.name} ({f * 100:.0f}%)' for n, f in pics) +
                    " — חילוץ: geometry_scan.py --extract-images %d" % idx))
    if tables:
        out.append(("info", "", f"{len(tables)} טבלה/ות חיות בשקף (תקין לפי #41)"))

    # ---- #2 רווח כפול בכל טקסט (כולל טבלאות/מסיחים/משוב)
    for n in text_nodes:
        for p in n.paras:
            for m in re.finditer(r"(?<=\S)( {2,3})(?=\S)", p["text"]):
                s = p["text"]
                out.append((2, "🔵", f'רווח כפול: "…{s[max(0, m.start() - 18): m.end() + 18]}…"'.replace("  ", "␣␣")))

    # ---- #44 רווח בין הטקסט המקדים לשאלה
    for n in on:
        if n.tag != "sp" or n in options or is_ui_label(n.text) or not n.paras:
            continue
        ps = n.paras
        ne = [i for i, p in enumerate(ps) if p["text"].strip()]
        if len(ne) < 2 or len(words(n.text)) < 10:
            continue
        qs = [i for i in ne if ps[i]["text"].strip().endswith("?")]
        if not qs or ne.index(qs[-1]) == 0:
            continue                      # אין משפט-שאלה (או שאין טקסט מקדים לפניו) — לשיפוט ידני
        li = qs[-1]
        last = ps[li]["text"].strip()
        if len(words(last)) > 22:
            continue
        pi = ne[ne.index(li) - 1]
        gap = (li - pi) > 1 or ps[pi]["spc_aft"] > 0 or ps[li]["spc_bef"] > 0 or ps[pi]["text"].endswith("\v") \
            or last.startswith("\v")
        if not gap:
            out.append((44, "🔵", f'אין רווח בין הטקסט המקדים לשאלה: "…{snippet(ps[pi]["text"].strip()[-26:], 30)}" ← '
                                  f'"{snippet(last, 40)}"'))

    # ---- #43 מילה בודדת בשורה
    for n in on:
        if n.tag != "sp" or not n.paras:
            continue
        ps = n.paras
        for i, p in enumerate(ps):
            t = p["text"]
            if "\v" in t:
                segs = [s.strip() for s in t.split("\v") if s.strip()]
                if len(segs) >= 2 and any(len(words(s)) == 1 for s in segs) and any(len(words(s)) >= 3 for s in segs):
                    one = [s for s in segs if len(words(s)) == 1][0]
                    if not one.endswith(":"):
                        out.append((43, "🔵", f'שבירת-שורה ידנית משאירה מילה בודדת בשורה: "{snippet(t, 50)}"'))
            elif len(words(t)) == 1 and not t.strip().endswith(":") and i > 0:
                prev = ps[i - 1]["text"].strip()
                if len(words(prev)) >= 3 and not prev.endswith((".", "?", "!", ":", ";")):
                    out.append((43, "🔵", f'מילה בודדת בפסקה נפרדת שממשיכה משפט: "{snippet(prev[-24:], 28)}" ⏎ "{t.strip()}"'))
        if n.wrap == "none" or n.box is None or re.search(r"שאלה \d+ מתוך \d+", n.text):
            continue          # תווית-התקדמות של התבנית ("שאלה 1 מתוך 5")
        for p in ps:
            t = p["text"].replace("\v", " ").strip()
            if len(words(t)) < 4:
                continue
            sz = p["sz"] or 18.0
            width = (n.box[2] - n.ins[0] - n.ins[2]) / EMU_PT
            if width < 40:
                continue
            lo = wrap_lines(t, sz, p["bold"], width * 0.96)
            hi = wrap_lines(t, sz, p["bold"], width * 1.04)
            if lo is None:
                break
            if len(lo) >= 2 and len(lo[-1]) == 1 and len(lo[-1][0]) <= 12:
                certainty = "סביר" if (len(hi) >= 2 and len(hi[-1]) == 1) else "גבולי"
                out.append((43, "🔵", f'מילה בודדת בשורה האחרונה ({certainty}, הערכת-גלישה {sz:g}pt): '
                                      f'"…{snippet(" ".join(lo[-2][-3:]), 24)} ⏎ {lo[-1][0]}"'))
    return out, stem


def main(argv):
    path = argv[0]
    only = None
    if "--slides" in argv:
        only = {int(x) for x in argv[argv.index("--slides") + 1].split(",")}
    prs = Presentation(path)
    W, H = prs.slide_width, prs.slide_height
    print("# סריקת תבנית וניקיון-הפקה (#37-#45 + #2 מורחב) — %d שקפים" % len(prs.slides))
    results, tally, infos = {}, Counter(), {}
    prev_stem = ""
    for i, slide in enumerate(prs.slides, 1):
        found, stem = scan_slide(slide, i, W, H, prev_stem)
        prev_stem = stem or ""
        if only and i not in only:
            continue
        real = [f for f in found if f[0] != "info"]
        info = [f for f in found if f[0] == "info"]
        if real:
            results[i] = real
            for c, _, _ in real:
                tally[c] += 1
        if info:
            infos[i] = info
    print("גופן למדידת-גלישה (#43): %s%s\n" % (FONT_USED["name"] or "לא נמצא (דילוג על הערכת-גלישה)",
                                              "" if (FONT_USED["name"] or "").startswith("Assistant")
                                              else "  ⚠ לא Assistant — הערכת-הגלישה גסה יותר"))
    print("## ממצאים (מועמדים) לפי שקף\n")
    for i in sorted(results):
        print("### שקף %d" % i)
        for c, sev, d in results[i]:
            print("- #%s %s: %s" % (c, sev, d))
        print()
    print("## מידע לקריאה שיפוטית (לא ממצא כשלעצמו)\n")
    for i in sorted(infos):
        for _, _, d in infos[i]:
            print("- שקף %d: %s" % (i, d))
    print("\n## סיכום לפי בדיקה")
    for c, n in sorted(tally.items(), key=lambda kv: str(kv[0])):
        print("- #%s: %d מופעים" % (c, n))
    print("\nשקפים עם ממצא: %d" % len(results))


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    sys.exit(main(sys.argv[1:]) or 0)
